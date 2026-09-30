"""`scripts/generate_charge_restricted_dummy.py`가 만든 37개 (정격용량, 바이어) 조합에 대응하는
충전 STEP 기준표(`ChargeProgramSpec`) 더미 워크북을 생성한다(`.docs/36` 추가 작업).

이 컴퓨터엔 실제 충전 프로그램 원본("★ 260428 AGM CF 충전 프로그램 3.xlsx")이 없어서(새로
클론한 환경), 그 파일이 없으면 `charge_program_deviation_pct`가 전부 NULL로 남아 1단 회귀가
"유효 표본 0건"으로 실패한다(직전 세션에서 실제로 재현·확인함). 이 스크립트는 그 문제를 근본
원인부터 없애기 위해, `backend/etl/charge_program.py`가 파싱하는 것과 동일한 구조(정격용량별
시트 + 바이어별 블록 + STEP/Total 행)의 더미 충전 프로그램 워크북을 만든다.

Total 충전량[Ah]은 이미 만들어 둔 더미 공정 데이터의 charge_ratio 분포(평균 112%, ±6%) 근방으로
맞추되, (정격용량, 바이어) 조합마다 108~116% 사이에서 결정적으로(해시 기반) 조금씩 다르게
잡는다 — 모든 조합에 똑같이 112%를 쓰면 charge_program_deviation_pct가 charge_ratio의 완전한
선형 변환이 돼(rated_capacity가 분자·분모에서 약분됨) VIF가 무한대로 발산하는 완전 다중공선성이
생긴다(1차 시도에서 실제로 재현·확인 — VIF 9999 캡 처리 자체는 정상 동작하지만, 회귀에서 이
컬럼이 독자적인 정보를 전혀 주지 못하는 건 더미 데이터 품질 문제다). 조합별로 기준치를 살짝
흩어 두면 charge_ratio와 강하게 상관되면서도 완전히 겹치지는 않는, 실제 충전 프로그램표처럼
바이어마다 조금씩 다른 기준을 가진 것과 같은 모양이 된다.

저장소 루트에서 실행: python scripts/generate_charge_program_spec_dummy.py
산출물: "충전바이어_더미_충전프로그램.xlsx" (저장소 루트)
"""
from __future__ import annotations

import hashlib

import openpyxl

OUT_SPEC = "충전바이어_더미_충전프로그램.xlsx"

# scripts/generate_charge_restricted_dummy.py의 BUYERS_BY_CAPACITY와 반드시 동일하게 유지 —
# 두 스크립트가 만드는 (정격용량, 바이어) 조합이 어긋나면 charge_program_deviation_pct가
# 다시 일부 로트에서만 NULL로 빠진다.
BUYERS_BY_CAPACITY: dict[float, list[str]] = {
    50.0: ["H", "K", "S", "Y"],
    60.0: ["H", "K", "S", "B", "C", "D", "G", "V", "Y", "Z"],
    70.0: ["H", "K", "S", "B", "C", "D", "G", "Y", "Z"],
    80.0: ["H", "K", "S", "B", "C", "D", "Y", "Z"],
    90.0: ["H", "K", "S"],
    105.0: ["H", "K", "S"],
}

CHARGE_RATIO_TARGET_PCT = 112.0  # generate_charge_restricted_dummy.py의 charge_ratio 평균과 일치

# backend/etl/charge_program.py: 원본 파일은 정격 90Ah 제품을 시트/블록 모두 "AGM92"로 표기하고
# 파서가 92->90으로 별칭 처리한다(_CAPACITY_ALIAS). 더미도 동일한 표기를 그대로 재현해야
# 실제 파서 경로(별칭 치환 포함)를 검증할 수 있다.
_SHEET_NAME_ALIAS = {90.0: "92"}
_TITLE_CAPACITY_ALIAS = {90.0: 92.0}


def _sheet_name(rated_capacity: float) -> str:
    return _SHEET_NAME_ALIAS.get(rated_capacity, str(int(rated_capacity)))


def _title_capacity(rated_capacity: float) -> int:
    return int(_TITLE_CAPACITY_ALIAS.get(rated_capacity, rated_capacity))


def _ratio_for_combo(rated_capacity: float, buyer: str) -> float:
    """조합별 108~116% 사이 결정적 지터 — charge_ratio와의 완전 선형 종속을 피한다(위 모듈
    docstring 참조)."""
    digest = hashlib.sha256(f"{rated_capacity}_{buyer}".encode()).hexdigest()
    unit = (int(digest[:8], 16) % 10_000) / 9_999  # 0.0~1.0
    return CHARGE_RATIO_TARGET_PCT - 4.0 + unit * 8.0  # 108.0~116.0


def build_workbook() -> openpyxl.Workbook:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    for rated_capacity, buyers in BUYERS_BY_CAPACITY.items():
        ws = wb.create_sheet(title=_sheet_name(rated_capacity))

        # 블록마다 5열(구분/전류[A]/시간[Hr]/충전량[Ah]/전기량[C])을 쓰고, 블록 사이는 빈 열 1개로
        # 구분한다(_find_block_starts가 title_row의 non-blank 셀 위치로 블록 시작을 잡기 때문에
        # 굳이 필요하진 않지만, 실제 원본 파일의 시각적 구조를 흉내내 실수를 줄인다).
        col_width = 5
        for block_idx, buyer in enumerate(buyers):
            col_start = block_idx * (col_width + 1) + 1  # 1-based (openpyxl)
            title = f"AGM{_title_capacity(rated_capacity)}_{buyer}"
            total_charge_ah = round(rated_capacity * _ratio_for_combo(rated_capacity, buyer) / 100, 2)

            ws.cell(row=1, column=col_start, value=title)
            ws.cell(row=2, column=col_start, value="구분")
            ws.cell(row=2, column=col_start + 1, value="전류[A]")
            ws.cell(row=2, column=col_start + 2, value="시간[Hr]")
            ws.cell(row=2, column=col_start + 3, value="충전량[Ah]")
            ws.cell(row=2, column=col_start + 4, value="전기량[C]")

            # STEP 1건(CHA1) — 이 더미의 목적은 charge_program_deviation_pct 계산용 Total
            # 충전량[Ah]을 제공하는 것이라 STEP 내역 자체의 현실성은 중요하지 않다(최소 구성만).
            ws.cell(row=3, column=col_start, value="CHA1")
            ws.cell(row=3, column=col_start + 1, value=round(rated_capacity * 0.1, 2))
            ws.cell(row=3, column=col_start + 2, value=11.0)
            ws.cell(row=3, column=col_start + 3, value=total_charge_ah)
            ws.cell(row=3, column=col_start + 4, value=round(total_charge_ah * 3600, 1))

            ws.cell(row=4, column=col_start, value="Total")
            ws.cell(row=4, column=col_start + 2, value=11.0)
            ws.cell(row=4, column=col_start + 3, value=total_charge_ah)
            ws.cell(row=4, column=col_start + 4, value=round(total_charge_ah * 3600, 1))

    return wb


def main() -> None:
    wb = build_workbook()
    wb.save(OUT_SPEC)

    from backend.etl.charge_program import parse_charge_program_workbook

    with open(OUT_SPEC, "rb") as f:
        rows = parse_charge_program_workbook(f.read())
    expected = sum(len(buyers) for buyers in BUYERS_BY_CAPACITY.values())
    assert len(rows) == expected, f"파싱된 행 수({len(rows)})가 기대값({expected})과 다릅니다"
    combos = {(r["rated_capacity"], r["buyer_code"]) for r in rows}
    expected_combos = {
        (cap, buyer) for cap, buyers in BUYERS_BY_CAPACITY.items() for buyer in buyers
    }
    assert combos == expected_combos, f"조합 불일치: {expected_combos ^ combos}"

    print(f"생성 완료: {OUT_SPEC} ({len(rows)}개 (정격용량, 바이어) 조합)")
    for r in sorted(rows, key=lambda r: (r["rated_capacity"], r["buyer_code"])):
        print(f"  AGM{int(r['rated_capacity'])}_{r['buyer_code']}: total_charge_ah={r['total_charge_ah']}")


if __name__ == "__main__":
    main()
